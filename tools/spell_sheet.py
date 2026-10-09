"""Make docs/spells.xlsx (and docs/spell-list.md) from the spells as the game shows them.

1. godot --headless --path . -s tools/dump_spells.gd -- <dump.json>   (every spell with its in-game card text)
2. python tools/spell_sheet.py <dump.json> [<git commit to compare with>]

The "Changed" column compares with data/spells.json at that commit (default: the last commit that touched
docs/spells.xlsx, i.e. the previous sheet). Spells gone since then are listed on a second sheet.
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
ELEMS = {"F": "Fire", "W": "Water", "A": "Air", "?": "Any"}
RARITY_ORDER = {"common": 0, "rare": 1, "legendary": 2}


def spell_type(s):
    if s["anti"]:
        return "Anti-spell"
    if s["power"]:
        return "Power"
    if s["fleeting"]:
        return "Fleeting"
    return "Starter" if s["starter"] else "Spell"


def git(*args):
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, encoding="utf-8").stdout


def changes(spells, base_commit):
    old = {s["id"]: s for s in json.loads(git("show", f"{base_commit}:data/spells.json"))["spells"]}
    now = {s["id"]: s for s in json.loads((ROOT / "data" / "spells.json").read_text(encoding="utf-8"))["spells"]}
    out = {}
    for sid, s in now.items():
        if sid not in old:
            out[sid] = "new"
            continue
        o = old[sid]
        what = []
        if o.get("name") != s.get("name"):
            what.append(f"renamed (was {o.get('name')})")
        if o.get("pattern") != s.get("pattern"):
            what.append(f"pattern (was {o.get('pattern') or '-'})")
        if o.get("rarity") != s.get("rarity"):
            what.append(f"rarity (was {o.get('rarity')})")
        if o.get("effects") != s.get("effects"):
            what.append("effect")
        if o.get("cooldown", 0) != s.get("cooldown", 0):
            what.append("cooldown")
        if bool(o.get("anti")) != bool(s.get("anti")) or bool(o.get("power")) != bool(s.get("power")):
            what.append("type")
        out[sid] = ", ".join(what)
    removed = [o for sid, o in old.items() if sid not in now]
    return out, removed


def style_header(ws, headers, widths, fill="4A5A3A"):
    thin = Side(style="thin", color="C8C8C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return border


def write_xlsx(spells, changed, removed, base_commit, path):
    font = Font(name="Arial", size=10)
    fills = {"F": "F5B7B1", "W": "AED6F1", "A": "D5F5E3"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Spells"
    headers = ["Name", "Pattern", "Length", "Rarity", "Category", "Fire", "Water", "Air", "Type", "Cooldown", "Effect (as on the card)", "Flavor", "Id", "Changed"]
    widths = [20, 22, 8, 11, 12, 6, 7, 6, 12, 10, 80, 44, 18, 30, 8]
    border = style_header(ws, headers, widths)
    rows = sorted(spells, key=lambda s: (len(s["pattern"]), s["pattern"], s["name"]))
    for r, s in enumerate(rows, 2):
        p = s["pattern"]
        vals = [s["name"], " → ".join(ELEMS.get(c, c) for c in p) if p else "(none: any chant)", f"=LEN(O{r})", s["rarity"].capitalize(),
                s["category"], p.count("F"), p.count("W"), p.count("A"), spell_type(s), s["cooldown"] or None, s["text"],
                s["flavor"], s["id"], changed.get(s["id"], "")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(wrap_text=c in (2, 11, 12, 14), vertical="top", horizontal="left" if c in (1, 2, 11, 12, 13, 14) else "center")
        ws.cell(r, 15, p).font = Font(name="Arial", size=8, color="999999")  # raw pattern, read by the Length formula
        for j, el in enumerate("FWA"):
            if p.count(el):
                ws.cell(r, 6 + j).fill = PatternFill("solid", fgColor=fills[el])
        if changed.get(s["id"]):
            ws.cell(r, 14).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.cell(1, 15, "code").font = Font(name="Arial", size=8, color="999999")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{len(rows) + 1}"
    note = len(rows) + 3
    ws.cell(note, 1, f"Effect text is taken from the game itself (SpellText.card_text). \"Changed\" compares with data/spells.json at commit {base_commit} (the previous sheet). Spells removed since then: see the Removed sheet.").font = Font(name="Arial", size=9, italic=True, color="666666")

    gone = wb.create_sheet("Removed")
    border = style_header(gone, ["Name", "Pattern", "Rarity", "Id"], [22, 14, 11, 20], fill="7A4A3A")
    for r, o in enumerate(sorted(removed, key=lambda o: o["name"]), 2):
        for c, v in enumerate([o["name"], o["pattern"], o["rarity"].capitalize(), o["id"]], 1):
            cell = gone.cell(r, c, v)
            cell.font = font
            cell.border = border
    if not removed:
        gone.cell(2, 1, "(none)").font = font
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def write_md(spells, path):
    lines = ["# Spell list", "", f"{len(spells)} spells, as the game shows them (made by tools/spell_sheet.py).", "",
             "| Name | Pattern | Rarity | Category | Type | Cooldown | Effect |", "|---|---|---|---|---|---|---|"]
    for s in sorted(spells, key=lambda s: (RARITY_ORDER[s["rarity"]], s["name"])):
        lines.append(f"| {s['name']} | {s['pattern'] or '-'} | {s['rarity'].capitalize()} | {s['category']} | {spell_type(s)} | {s['cooldown'] or ''} | {s['text']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    spells = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    base = sys.argv[2] if len(sys.argv) > 2 else git("log", "-1", "--format=%h", "--", "docs/spells.xlsx").strip()
    changed, removed = changes(spells, base)
    write_xlsx(spells, changed, removed, base, ROOT / "docs" / "spells.xlsx")
    write_md(spells, ROOT / "docs" / "spell-list.md")
    print(f"OK: {len(spells)} spells, {sum(1 for v in changed.values() if v)} changed or new since {base}, {len(removed)} removed")


if __name__ == "__main__":
    main()
