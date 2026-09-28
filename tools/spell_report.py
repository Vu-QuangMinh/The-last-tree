"""Validate data/spells.json (chant patterns, v3) and render docs/spell-list.md + docs/spells.xlsx.

Usage: python tools/spell_report.py
Exits non-zero if any rule is violated.
"""
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ELEMS = {"F": "Fire", "W": "Water", "A": "Air"}
ICON = {"F": "🔥", "W": "💧", "A": "🌪️"}
RARITIES = ("common", "rare", "legendary")
TARGETED = {"steal", "pluck", "move", "redirect", "strike", "burn", "poison", "siphon", "execute", "insert", "curse", "weak", "freeze", "expose", "ethereal", "rotate", "swap", "convert", "purge", "shatter", "stoke"}
OPS = {
    "strike": {"from", "n", "target", "min_hp", "bonus_if"},
    "burn": {"n", "target"}, "poison": {"n", "target"}, "stoke": {"target"}, "weak": {"turns", "target"}, "freeze": {"turns", "target"},
    "expose": {"turns", "target"}, "ethereal": {"target"}, "shield": {"n"}, "heal": {"n"}, "aegis": {"n"},
    "thorns": {"n"}, "draw": {"n", "when", "el", "temp"}, "rotate": {"target", "dir"}, "swap": {"target"},
    "convert": {"pos", "from", "to", "target"}, "purge": {"el", "n", "target"}, "shatter": {"target"},
    "amplify": {"n"}, "echo": {"n"}, "retain": {"n"}, "overload": {"n"}, "cleanse": {"what"},
    "siphon": {"n", "target"}, "execute": {"max", "target"}, "insert": {"el", "pos", "target"},
    "transmute": {"n", "to"}, "sacrifice": {"hp"}, "copy_last": set(),
    "passive": {"key", "n"}, "move": {"n", "target"}, "pluck": {"n", "target"}, "steal": {"n", "el", "target"}, "infuse": {"el"}, "duplicate": {"times"}, "summon_spells": {"n"}, "redirect": {"target"}, "curse": {"key", "target"}, "each_turn": {"effects"},
}


def who(t):
    return {"target": "the target", "two": "2 different enemies (you pick both)", "all": "every enemy", "random": "a random enemy", "self": "you"}.get(t, t)


def turns(n):
    return f"{n} turn" + ("" if n == 1 else "s")


def describe_op(e):
    op = e["op"]
    t = who(e.get("target", "target"))
    if op == "strike":
        last = e["from"] == "right"
        side = "last" if last else "first"
        if e["n"] == 1:
            s = f"remove the {side} ({'rightmost' if last else 'leftmost'}) Essence of {t}"
        else:
            s = f"remove the {side} {e['n']} Essence of {t}, {'right to left' if last else 'left to right'}"
        if "min_hp" in e:
            s += f" (only if it has {e['min_hp']}+ Essence)"
        if "bonus_if" in e:
            s += " (+1 if it is burning)"
        return s
    if op == "burn":
        return f"Burn {e['n']} on {t}"
    if op == "poison":
        return f"Poison {e['n']} on {t}"
    if op == "stoke":
        return f"double the Burn on {t}"
    if op == "weak":
        return f"Weaken {t} for {turns(e['turns'])}"
    if op == "freeze":
        return f"Freeze {t} for {turns(e['turns'])}"
    if op == "expose":
        return f"Expose {t} for {turns(e['turns'])}"
    if op == "ethereal":
        return "you become Ethereal this turn" if e["target"] == "self" else f"make {t} Ethereal (your spells remove double from it this turn)"
    if op == "shield":
        return f"gain {e['n']} Shield"
    if op == "heal":
        return f"heal {e['n']}"
    if op == "aegis":
        return f"gain Aegis ({e['n']} hit)"
    if op == "thorns":
        return f"Thorns {e['n']} this turn"
    if op == "draw":
        el = "random element" if e["el"] == "random" else ELEMS[e["el"]]
        el = ("conjured " if e.get("temp") else "") + el
        return f"+{e['n']} {el}{'s' if e['n'] > 1 and e['el'] == 'random' else ''} {'next turn' if e['when'] == 'next' else 'now'}"
    if op == "rotate":
        if e.get("dir") == "right":
            return f"move the last element of {t} to the front"
        return f"move the first element of {t} to the end"
    if op == "steal":
        if e.get("el", "any") == "any":
            return f"Steal {e['n']} element{'s' if e['n'] > 1 else ''} of your choice from {t} (they go to your elements)"
        return f"Steal up to {e['n']} {ELEMS[e['el']]} from {t} (they go to your elements)"
    if op == "infuse":
        return f"put a {ELEMS[e['el']]} anywhere you like in this turn's chant"
    if op == "duplicate":
        return f"pick an element in this turn's chant: it becomes {e['times']} in a row"
    if op == "summon_spells":
        return f"add {e['n']} random spells from your spellbook to your active row for this fight"
    if op == "pluck":
        return f"Targeted: remove {e['n']} Essence of your choice from {t}"
    if op == "move":
        return f"move {e['n']} element{'s' if e['n'] > 1 else ''} of {t} to any position you choose"
    if op == "redirect":
        return f"redirect the intent of {t}: its attacks hit the enemy you choose (it can be itself), and anything aimed at you fizzles"
    if op == "swap":
        return f"swap the first two elements of {t}"
    if op == "convert":
        where = {"first": "first element", "first2": "first two elements", "last": "last element", "all": f"every {ELEMS[e['from']]}" if "from" in e else "every element"}[e["pos"]]
        return f"turn {'' if e['pos'] == 'all' else 'the '}{where} of {t} into {ELEMS[e['to']]}"
    if op == "purge":
        return f"remove up to {e['n']} {ELEMS[e['el']]} from {t}"
    if op == "shatter":
        return f"remove all Armour from {t}"
    if op == "amplify":
        return f"Amplify {e['n']}: when the chant is Released, every enemy it hits also loses its last {e['n']} Essence, right to left"
    if op == "echo":
        return "Echo: the chant is Released twice this turn"
    if op == "retain":
        return f"after casting, {e['n']} of the chant's elements return to your stock"
    if op == "overload":
        return f"Overload {e['n']}: {e['n']} fewer element next turn"
    if op == "cleanse":
        return {"all": "remove all your debuffs (frozen elements, Blind, Confuse, Bleed, Silence)",
                "blind": "cure Blind", "bleed": "cure Bleed", "confuse": "cure Confuse",
                "silence": "cure Silence on all your spells", "frozen": "thaw your frozen elements",
                "lock": "break one Lock on your spells"}[e.get("what", "all")]
    if op == "siphon":
        return f"take the first {e['n']} element{'s' if e['n'] > 1 else ''} of {t}; you get {'them' if e['n'] > 1 else 'it'} next turn as conjured"
    if op == "execute":
        return f"if {t} has {e['max']} or fewer elements, destroy it"
    if op == "insert":
        return f"put a {ELEMS[e['el']]} at the front of {t}"
    if op == "transmute":
        return f"turn {e['n']} of your stored elements into {ELEMS[e['to']]}"
    if op == "sacrifice":
        return f"lose {e['hp']} HP"
    if op == "copy_last":
        return "repeat the previous spell that fired this cast"
    if op == "passive":
        return {"burn_bonus": f"your Burn applies +{e['n']} stack",
                "thorns": f"Thorns {e['n']} for the rest of the fight",
                "strike_poison": f"every enemy your Release hits gets Poison {e['n']}",
                "strike_burn": f"every enemy your Release hits gets Burn {e['n']}",
                "echo_first": "the first spell you trigger each turn fires twice",
                "attune": "choose an element: one of your draws each turn is always that element",
                "chant_slots": f"your chant line gets +{e['n']} slots",
                "strike_bonus": f"your spells that remove elements remove +{e['n']}"}[e["key"]]
    if op == "curse":
        return {"no_mend": f"{t} can never heal or regrow elements",
                "exposed": f"{t} is permanently Exposed",
                "weak25": "every enemy permanently deals 25% less damage"}[e["key"]]
    if op == "each_turn":
        return "at the start of each of your turns: " + "; ".join(describe_op(x) for x in e["effects"])
    return op


DAMAGE_OPS = {"steal", "pluck", "strike", "burn", "poison", "purge", "execute", "siphon", "amplify", "echo", "stoke", "expose"}
DEFENSE_OPS = {"shield", "heal", "aegis", "thorns", "weak", "freeze", "cleanse", "redirect"}
DAMAGE_KEYS = {"burn_bonus", "strike_poison", "strike_burn", "strike_bonus", "exposed"}
DEFENSE_KEYS = {"thorns", "weak25"}


def kind_of(effects):
    flat = []
    for e in effects:
        flat += e["effects"] if e["op"] == "each_turn" else [e]
    if any(e["op"] in DAMAGE_OPS or (e["op"] in ("passive", "curse") and e["key"] in DAMAGE_KEYS) for e in flat):
        return "Offensive"
    if any(e["op"] in DEFENSE_OPS or (e["op"] in ("passive", "curse") and e["key"] in DEFENSE_KEYS) or (e["op"] == "ethereal" and e["target"] == "self") for e in flat):
        return "Defensive"
    return "Utility"


def describe(s):
    text = "; ".join(describe_op(e) for e in s["effects"])
    text = text[0].upper() + text[1:] + "."
    if s.get("power"):
        text = "**Power** (leaves your active row for the rest of the fight): " + text
    return text


def _check_effects(s, effects, errors):
    for e in effects:
        if e.get("op") not in OPS:
            errors.append(f"{s['id']}: unknown op {e.get('op')}")
            continue
        extra = set(e) - {"op"} - OPS[e["op"]]
        if extra:
            errors.append(f"{s['id']}: {e['op']} has unknown keys {extra}")
        if e["op"] in TARGETED and e.get("target") not in ("target", "two", "all", "random", "self"):
            errors.append(f"{s['id']}: {e['op']} needs a target")
        if e["op"] in ("passive", "curse", "each_turn", "summon_spells") and not s.get("power"):
            errors.append(f"{s['id']}: {e['op']} only belongs on a Power")
        if e["op"] == "each_turn":
            _check_effects(s, e["effects"], errors)


def validate(data):
    errors = []
    spells = data["spells"]
    ids = Counter(s["id"] for s in spells)
    errors += [f"duplicate id: {i}" for i, n in ids.items() if n > 1]
    for s in spells:
        p = s["pattern"]
        if not p or any(c not in ELEMS for c in p):
            errors.append(f"{s['id']}: bad pattern {p}")
        if not 1 <= len(p) <= 5:
            errors.append(f"{s['id']}: pattern length {len(p)}")
        _check_effects(s, s["effects"], errors)
    lengths = Counter(len(s["pattern"]) for s in spells)
    # bell curve: 2 most common, then 3 and 1, 4 rare, 5 rarest
    if not (lengths[2] >= lengths[3] >= lengths[4] >= lengths[5] and lengths[2] >= lengths[1] >= lengths[5]):
        errors.append(f"pattern lengths should be bell-shaped around 2: {dict(sorted(lengths.items()))}")
    for sp in spells:
        if sp.get("rarity") not in RARITIES:
            errors.append(f"{sp['id']}: rarity must be one of {RARITIES}")
    starters = [s["id"] for s in spells if s.get("starter")]
    if sorted(starters) != sorted(["fire_ball", "water_wall", "tailwind"]):
        errors.append(f"starters must be Fire Ball, Water Wall, Tailwind (got {starters})")
    return errors


def fmt_pattern(p):
    return " ".join(ICON[c] for c in p)


def render(data):
    spells = data["spells"]
    lines = [
        "# The Last Tree — Spell List (chant patterns)",
        "",
        "Generated from `data/spells.json` by `tools/spell_report.py`. Edit the JSON, then re-run the script.",
        "",
        f"**{len(spells)} spells.** Chant, then every spell whose pattern appears in the chant comes alive: one charge per "
        "separate match, but never more than its pattern's length. Cast them in any order; then the chant is Released "
        "and each enemy loses the longest start of its Essence found in it. ★ = starter.",
        "",
        "**Keywords**",
        "- **Burn N:** at the start of its turn, the enemy loses its **first** (leftmost) element, then Burn goes down by 1.",
        "- **Poison N:** like Burn, but the enemy loses its **last** (rightmost) element.",
        "- **Weaken:** the enemy deals 50% less damage.",
        "- **Freeze:** the enemy skips its next action.",
        "- **Expose:** whenever your Release hits this enemy, it also loses its last (rightmost) element.",
        "- **Ethereal (you):** take no damage from attacks this turn, but double damage from effects. "
        "**Ethereal (enemy):** your spells remove double from it this turn; enemies that go Ethereal as their intent "
        "can't be hit by your next Release.",
        "- **Shield:** blocks damage until your next turn. **Aegis:** blocks one hit completely.",
        "- **Thorns:** enemies that attack you lose their rightmost element.",
        "- **Armour** (enemy ability): an armoured element can't be removed this turn, but it still counts for the chant's match.",
        "- **Power:** cast once, then it leaves your active row (its slot stays empty) and its effect lasts the whole fight.",
        "- **Conjured** elements arrive next turn and vanish at the end of that turn if unused.",
        "- **Siphon:** take elements off an enemy's Essence. **Execute:** destroy an enemy that is small enough.",
        "- **Amplify:** enemies your chant hit lose extra elements. **Echo:** the chant strikes again after your spells. "
        "**Overload:** fewer elements next turn.",
        "",
    ]
    for n in range(1, 6):
        group = [s for s in spells if len(s["pattern"]) == n]
        lines += [f"## {n}-element patterns ({len(group)})", "",
                  "| Name | Pattern | Rarity | Category | Effect | Flavor |", "|---|---|---|---|---|---|"]
        for s in sorted(group, key=lambda s: (s["pattern"], s["name"])):
            star = " ★" if s.get("starter") else ""
            lines.append(f"| **{s['name']}**{star} | {fmt_pattern(s['pattern'])} | {s['rarity'].capitalize()} | {kind_of(s['effects'])} | {describe(s)} | *{s.get('flavor', '')}* |")
        lines.append("")
    return "\n".join(lines)


def write_xlsx(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    font = Font(name="Arial", size=10)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="C8C8C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {"F": "F5B7B1", "W": "AED6F1", "A": "D5F5E3"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Spells"
    headers = ["Name", "Pattern", "Length", "Rarity", "Category", "Fire", "Water", "Air", "Starter", "Effect", "Flavor"]
    widths = [20, 22, 8, 11, 12, 6, 7, 6, 8, 90, 44]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = head
        cell.fill = PatternFill("solid", fgColor="4A5A3A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    rows = sorted(data["spells"], key=lambda s: (len(s["pattern"]), s["pattern"], s["name"]))
    for r, s in enumerate(rows, 2):
        p = s["pattern"]
        vals = [s["name"], " → ".join(ELEMS[c] for c in p), f"=LEN(L{r})", s["rarity"].capitalize(), kind_of(s["effects"]),
                p.count("F"), p.count("W"), p.count("A"), "yes" if s.get("starter") else "", describe(s), s.get("flavor", "")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(wrap_text=c in (2, 10, 11), vertical="top", horizontal="left" if c in (1, 2, 10, 11) else "center")
        ws.cell(r, 12, p).font = Font(name="Arial", size=8, color="999999")  # raw pattern, used by the Length formula
        for j, el in enumerate("FWA"):
            if p.count(el):
                ws.cell(r, 6 + j).fill = PatternFill("solid", fgColor=fills[el])
    for c, w in enumerate(widths + [8], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.cell(1, 12, "code").font = Font(name="Arial", size=8, color="999999")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{len(rows) + 1}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main():
    data = json.loads((ROOT / "data" / "spells.json").read_text(encoding="utf-8"))
    errors = validate(data)
    if errors:
        print("\n".join(errors))
        sys.exit(1)
    out = ROOT / "docs" / "spell-list.md"
    out.write_text(render(data), encoding="utf-8")
    counts = Counter(c for s in data["spells"] for c in s["pattern"])
    print("Element slots:", {ELEMS[k]: counts[k] for k in ELEMS})
    print("Categories:", dict(Counter(kind_of(s["effects"]) for s in data["spells"])))
    xlsx = ROOT / "docs" / "spells.xlsx"
    try:
        write_xlsx(data, xlsx)
        print(f"OK: {len(data['spells'])} spells valid -> {out}, {xlsx}")
    except PermissionError:
        print(f"OK: {len(data['spells'])} spells valid -> {out}")
        print(f"WARNING: {xlsx} is locked (open in Excel?) - close it and re-run to refresh the spreadsheet.")


if __name__ == "__main__":
    main()
